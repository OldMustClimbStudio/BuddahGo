"""In-memory model of the design-table workbooks plus the xlsx / CSV writers and an xlsx reader.

Shared by build_tables.py (baseline generation) and migrate_v2.py (in-place restructuring) so both paths
produce byte-identical CSV mirrors and the same workbook styling.

Model: Model.workbooks[wb_name] -> {sheet_name: Sheet}, insertion-ordered.
Sheet.cols  = [(var, type, group, comment)]
Sheet.rows  = [dict var -> value]; a row may carry '_marker' ('##' comments the row out in column A).
Sheet.record / source / notes feed Global/__meta__.
"""
import csv, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
DOC_FILL = PatternFill('solid', fgColor='F2F2F2')
KEY_FONT = Font(bold=True)
META_SHEET = '__meta__'
ENUM_SHEET = '__enums__'


def fmt(x):
    if isinstance(x, bool):
        return 'true' if x else 'false'
    if isinstance(x, float):
        if x == int(x) and abs(x) < 1e9:
            return str(int(x))
        return repr(round(x, 6))
    if x is None:
        return ''
    return str(x)


class Sheet:
    def __init__(self, workbook, name, cols, rows, record='(新增)', source='', notes=''):
        self.workbook = workbook
        self.name = name
        self.cols = list(cols)
        self.rows = list(rows)
        self.record = record
        self.source = source
        self.notes = notes

    # --- column helpers -------------------------------------------------------------------------
    @property
    def vars(self):
        return [c[0] for c in self.cols]

    def col(self, var):
        for c in self.cols:
            if c[0] == var:
                return c
        raise KeyError(f'{self.workbook}/{self.name} has no column {var}')

    def has(self, var):
        return any(c[0] == var for c in self.cols)

    def set_col(self, var, typ=None, grp=None, com=None):
        """Change type / group / comment of an existing column (None keeps the current value)."""
        for i, (v, t, g, c) in enumerate(self.cols):
            if v == var:
                self.cols[i] = (v, typ if typ is not None else t, grp if grp is not None else g, com if com is not None else c)
                return
        raise KeyError(var)

    def rename_col(self, old, new):
        for i, (v, t, g, c) in enumerate(self.cols):
            if v == old:
                self.cols[i] = (new, t, g, c)
        for r in self.rows:
            if old in r:
                r[new] = r.pop(old)

    def drop_col(self, var):
        self.cols = [c for c in self.cols if c[0] != var]
        for r in self.rows:
            r.pop(var, None)

    def insert_col(self, after, var, typ, com='', grp=''):
        idx = [c[0] for c in self.cols].index(after) + 1 if after else 0
        self.cols.insert(idx, (var, typ, grp, com))

    def row(self, **key):
        for r in self.rows:
            if all(str(r.get(k, '')) == str(v) for k, v in key.items()):
                return r
        return None

    def value(self, var, **key):
        r = self.row(**key)
        return r.get(var, '') if r else ''


class Model:
    def __init__(self):
        self.workbooks = {}

    def add(self, sheet):
        self.workbooks.setdefault(sheet.workbook, {})[sheet.name] = sheet
        return sheet

    def get(self, wb, name):
        return self.workbooks[wb][name]

    def find(self, wb, name):
        return self.workbooks.get(wb, {}).get(name)

    def remove(self, wb, name):
        return self.workbooks[wb].pop(name)

    def sheets(self):
        for wb in self.workbooks.values():
            for s in wb.values():
                yield s

    def insert_after(self, wb, after, sheet):
        """Add a sheet right after another one (keeps a readable workbook order)."""
        d = self.workbooks.setdefault(wb, {})
        items = list(d.items())
        d.clear()
        placed = False
        for k, v in items:
            d[k] = v
            if k == after:
                d[sheet.name] = sheet
                placed = True
        if not placed:
            d[sheet.name] = sheet
        return sheet

    def rebuild_meta(self):
        rows = [{'workbook': s.workbook, 'sheet': s.name, 'recordClass': s.record, 'source': s.source, 'notes': s.notes}
                for s in self.sheets() if s.name != META_SHEET]
        meta = Sheet('Global', META_SHEET, [('workbook', 'string#key', '', ''), ('sheet', 'string#key', '', ''), ('recordClass', 'string', '', ''),
                                            ('source', 'string', '', '实际值来源'), ('notes', 'string', '', '')], rows, '元信息（不导出）', '')
        self.workbooks.setdefault('Global', {}).pop(META_SHEET, None)
        self.workbooks['Global'][META_SHEET] = meta
        return meta


# --- reading ----------------------------------------------------------------------------------------
def read_workbook(path, workbook_name, meta_lookup=None):
    """Load every ##-headed sheet of an xlsx into Sheet objects. Values keep openpyxl types (int/float/str)."""
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        a1 = ws.cell(1, 1).value
        if not (isinstance(a1, str) and a1.startswith('##')):
            continue
        header = {}
        first_data = None
        for r in range(1, ws.max_row + 1):
            marker = ws.cell(r, 1).value
            if isinstance(marker, str) and marker in ('##var', '##type', '##group', '##comment') and marker not in header:
                header[marker] = r
                continue
            first_data = r
            break
        var_r, type_r = header['##var'], header['##type']
        grp_r, com_r = header.get('##group'), header.get('##comment')
        cols = []
        col_index = []
        for c in range(2, ws.max_column + 1):
            var = ws.cell(var_r, c).value
            if var is None or str(var).strip() == '':
                continue
            cols.append((str(var).strip(), _s(ws.cell(type_r, c).value), _s(ws.cell(grp_r, c).value) if grp_r else '', _s(ws.cell(com_r, c).value) if com_r else ''))
            col_index.append(c)
        rows = []
        if first_data is not None:
            for r in range(first_data, ws.max_row + 1):
                marker = ws.cell(r, 1).value
                vals = [ws.cell(r, c).value for c in col_index]
                if all(v is None or str(v).strip() == '' for v in vals) and not (isinstance(marker, str) and marker.startswith('##')):
                    continue
                row = {cols[i][0]: ('' if v is None else v) for i, v in enumerate(vals)}
                if isinstance(marker, str) and marker.startswith('##'):
                    row['_marker'] = marker
                rows.append(row)
        rec, src, notes = ('', '', '')
        if meta_lookup and (workbook_name, ws.title) in meta_lookup:
            rec, src, notes = meta_lookup[(workbook_name, ws.title)]
        sheets.append(Sheet(workbook_name, ws.title, cols, rows, rec, src, notes))
    return sheets


def _s(v):
    return '' if v is None else str(v)


def read_tables(tables_dir):
    """Read all <Domain>/<Domain>.xlsx under tables_dir into a Model, attaching __meta__ info."""
    model = Model()
    names = sorted(d for d in os.listdir(tables_dir) if os.path.isfile(os.path.join(tables_dir, d, d + '.xlsx')))
    order = ['Skill', 'Movement', 'AI', 'Match', 'UI', 'Global']
    names.sort(key=lambda n: order.index(n) if n in order else 99)
    meta_lookup = {}
    if 'Global' in names:
        for s in read_workbook(os.path.join(tables_dir, 'Global', 'Global.xlsx'), 'Global'):
            if s.name == META_SHEET:
                for r in s.rows:
                    meta_lookup[(r['workbook'], r['sheet'])] = (_s(r.get('recordClass')), _s(r.get('source')), _s(r.get('notes')))
    for n in names:
        for s in read_workbook(os.path.join(tables_dir, n, n + '.xlsx'), n, meta_lookup):
            if s.name == META_SHEET:
                continue
            model.add(s)
    return model


# --- writing ----------------------------------------------------------------------------------------
def _cell_value(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'true' if v else 'false'
    if isinstance(v, float):
        s = fmt(v)
        return float(s) if s.replace('-', '').replace('.', '').isdigit() else v
    return v


def write_sheet(wb, sheet):
    if wb.active and wb.active.title == 'Sheet':
        ws = wb.active
        ws.title = sheet.name
    else:
        ws = wb.create_sheet(sheet.name)
    ws['A1'] = '##var'; ws['A2'] = '##type'; ws['A3'] = '##group'; ws['A4'] = '##comment'
    for j, (var, typ, grp, com) in enumerate(sheet.cols, start=2):
        ws.cell(1, j, var).font = KEY_FONT if '#key' in typ else Font()
        ws.cell(2, j, typ); ws.cell(3, j, grp); ws.cell(4, j, com)
        fill = DOC_FILL if grp == 'doc' else HEAD_FILL
        for r in range(1, 5):
            ws.cell(r, j).fill = fill
        ws.cell(4, j).alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(48, len(var) + 4))
    ws.column_dimensions['A'].width = 11
    for i, row in enumerate(sheet.rows, start=5):
        if row.get('_marker'):
            ws.cell(i, 1, row['_marker'])
        for j, (var, typ, grp, com) in enumerate(sheet.cols, start=2):
            ws.cell(i, j, _cell_value(row.get(var, '')))
    ws.freeze_panes = 'C5'
    return ws


def write_csv(out_root, sheet):
    """CSV mirror in the exporter's format: header rows, then data rows; commented-out rows are omitted."""
    d = os.path.join(out_root, '_Export', sheet.workbook)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, f'{sheet.name}.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['##var'] + [c[0] for c in sheet.cols]); w.writerow(['##type'] + [c[1] for c in sheet.cols])
        w.writerow(['##group'] + [c[2] for c in sheet.cols]); w.writerow(['##comment'] + [c[3] for c in sheet.cols])
        for row in sheet.rows:
            if row.get('_marker'):
                continue
            w.writerow([''] + [_csv_value(row.get(c[0], '')) for c in sheet.cols])


def _csv_value(v):
    if v is None:
        return ''
    return v if isinstance(v, str) else fmt(v)


def write_workbook_xlsx(model, out_root, name, path=None):
    """Write one workbook to <out_root>/<name>/<name>.xlsx (or an explicit path)."""
    wb = Workbook()
    for s in model.workbooks[name].values():
        write_sheet(wb, s)
    d = os.path.join(out_root, name)
    os.makedirs(d, exist_ok=True)
    target = path or os.path.join(d, f'{name}.xlsx')
    wb.save(target)
    return target


def write_workbook_csvs(model, out_root, name):
    """Write the CSV mirror of one workbook and delete mirrors of sheets that no longer exist."""
    sheets = model.workbooks[name]
    export_dir = os.path.join(out_root, '_Export', name)
    if os.path.isdir(export_dir):
        for f in os.listdir(export_dir):
            if f.endswith('.csv') and f[:-4] not in sheets:
                os.remove(os.path.join(export_dir, f))
                meta = os.path.join(export_dir, f + '.meta')
                if os.path.exists(meta):
                    os.remove(meta)
    for s in sheets.values():
        write_csv(out_root, s)


def write_model(model, out_root, workbooks=None):
    """Write <wb>/<wb>.xlsx and _Export/<wb>/*.csv for the given workbooks (None = all). Rebuilds Global/__meta__."""
    model.rebuild_meta()
    written = []
    for name in model.workbooks:
        if workbooks is not None and name not in workbooks:
            continue
        write_workbook_xlsx(model, out_root, name)
        write_workbook_csvs(model, out_root, name)
        written.append((name, list(model.workbooks[name])))
    return written
